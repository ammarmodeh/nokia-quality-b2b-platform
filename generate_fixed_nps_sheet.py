"""
Fix and generate the complete NPS Analysis and Feedback Excel Workbook
with corrected Ownership Distribution (100.0% sum across total issues).
"""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_fixed_nps_workbook():
    wb = openpyxl.Workbook()
    
    # Setup styling constants
    NAVY_HEADER = "002060"
    DARK_GREEN = "006100"
    LIGHT_GREEN = "C6EFCE"
    DARK_BLUE = "002060"
    ACC_YES_BLUE = "1F4E79"
    ACC_NO_RED = "C00000"
    DETRACTOR_RED = "C00000"
    NEUTRAL_ORANGE = "ED7D31"
    WHITE = "FFFFFF"
    BORDER_COLOR = "D9D9D9"

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )
    thick_bottom = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='medium', color="000000")
    )

    # -------------------------------------------------------------
    # SHEET: Ownership Distribution
    # -------------------------------------------------------------
    ws = wb.active
    ws.title = "Ownership Distribution"
    ws.views.sheetView[0].showGridLines = True

    # Main Banner
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = "OWNERSHIP DISTRIBUTION - Period: 01/09/2026 - 30/09/2026"
    title_cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    title_cell.font = Font(name="Calibri", size=14, bold=True, color=WHITE)
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28

    # Section 1 Header: OWNER DISTRIBUTION SUMMARY
    ws.merge_cells("A2:K2")
    s1_cell = ws["A2"]
    s1_cell.value = "OWNER DISTRIBUTION SUMMARY"
    s1_cell.fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    s1_cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    s1_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 22

    # Level 1 Headers (Merged)
    headers_l1 = [
        ("A3:A4", "Owner / Entity", "002060"),
        ("B3:C3", "TOTAL ISSUES IDENTIFIED (100% DISTRIBUTION)", "1F4E79"),
        ("D3:E3", "ACCOUNTABLE (YES)", "1F4E79"),
        ("F3:G3", "NON-ACCOUNTABLE (NO)", "C00000"),
        ("H3:I3", "DETRACTORS", "C00000"),
        ("J3:K3", "NEUTRALS", "ED7D31"),
    ]

    for rng, text, bg in headers_l1:
        ws.merge_cells(rng)
        first_cell = ws[rng.split(":")[0]]
        first_cell.value = text
        first_cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
        first_cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        first_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Level 2 Sub-headers
    sub_headers = {
        "B4": ("Count", "1F4E79"),
        "C4": ("Percentage", "1F4E79"),
        "D4": ("Team Acc. Yes", "1F4E79"),
        "E4": ("Yes %", "1F4E79"),
        "F4": ("Team Acc. No", "C00000"),
        "G4": ("No %", "C00000"),
        "H4": ("Detractors", "C00000"),
        "I4": ("Detractors %", "C00000"),
        "J4": ("Neutrals", "ED7D31"),
        "K4": ("Neutrals %", "ED7D31"),
    }
    for cell_id, (txt, bg) in sub_headers.items():
        c = ws[cell_id]
        c.value = txt
        c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
        c.font = Font(name="Calibri", size=9.5, bold=True, color=WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 22

    # Data Rows
    # [Owner, IssuesCount, AccYes, AccNo, Detractors, Neutrals]
    data_summary = [
        ("Reach", 6, 4, 2, 4, 2),
        ("Customer", 3, 0, 3, 2, 1),
        ("Not specified", 2, 0, 2, 0, 2),
        ("OJO", 2, 0, 2, 2, 0),
    ]

    for idx, row in enumerate(data_summary, start=5):
        owner, issues, acc_y, acc_n, detr, neutr = row
        
        ws[f"A{idx}"] = owner
        ws[f"A{idx}"].font = Font(name="Calibri", size=10, bold=True)
        ws[f"A{idx}"].alignment = Alignment(horizontal="left", vertical="center")

        # Column B & C: Total Issues Identified (Primary Distribution = 100%)
        ws[f"B{idx}"] = issues
        ws[f"B{idx}"].font = Font(name="Calibri", size=10, bold=True, color="002060")
        ws[f"B{idx}"].alignment = Alignment(horizontal="center", vertical="center")
        
        ws[f"C{idx}"] = f"=B{idx}/$B$9"
        ws[f"C{idx}"].number_format = "0.0%"
        ws[f"C{idx}"].alignment = Alignment(horizontal="center", vertical="center")

        # Column D & E: Accountable Yes
        ws[f"D{idx}"] = acc_y
        ws[f"D{idx}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"E{idx}"] = f"=IF(B{idx}>0, D{idx}/B{idx}, 0)"
        ws[f"E{idx}"].number_format = "0.0%"
        ws[f"E{idx}"].alignment = Alignment(horizontal="center", vertical="center")

        # Column F & G: Non-Accountable No
        ws[f"F{idx}"] = acc_n
        ws[f"F{idx}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"G{idx}"] = f"=IF(B{idx}>0, F{idx}/B{idx}, 0)"
        ws[f"G{idx}"].number_format = "0.0%"
        ws[f"G{idx}"].alignment = Alignment(horizontal="center", vertical="center")

        # Column H & I: Detractors
        ws[f"H{idx}"] = detr
        ws[f"H{idx}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"I{idx}"] = f"=IF(B{idx}>0, H{idx}/B{idx}, 0)"
        ws[f"I{idx}"].number_format = "0.0%"
        ws[f"I{idx}"].alignment = Alignment(horizontal="center", vertical="center")

        # Column J & K: Neutrals
        ws[f"J{idx}"] = neutr
        ws[f"J{idx}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"K{idx}"] = f"=IF(B{idx}>0, J{idx}/B{idx}, 0)"
        ws[f"K{idx}"].number_format = "0.0%"
        ws[f"K{idx}"].alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[idx].height = 20
        for col_letter in ["A","B","C","D","E","F","G","H","I","J","K"]:
            ws[f"{col_letter}{idx}"].border = thin_border

    # TOTAL ROW (Row 9)
    ws["A9"] = "TOTAL"
    ws["A9"].font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws["A9"].fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    ws["A9"].alignment = Alignment(horizontal="left", vertical="center")

    ws["B9"] = "=SUM(B5:B8)"
    ws["B9"].font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws["B9"].fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    ws["B9"].alignment = Alignment(horizontal="center", vertical="center")

    ws["C9"] = "=SUM(C5:C8)"
    ws["C9"].number_format = "0.0%"
    ws["C9"].font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws["C9"].fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    ws["C9"].alignment = Alignment(horizontal="center", vertical="center")

    ws["D9"] = "=SUM(D5:D8)"
    ws["D9"].font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws["D9"].fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    ws["D9"].alignment = Alignment(horizontal="center", vertical="center")

    ws["E9"] = "=D9/B9"
    ws["E9"].number_format = "0.0%"
    ws["E9"].font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws["E9"].fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    ws["E9"].alignment = Alignment(horizontal="center", vertical="center")

    ws["F9"] = "=SUM(F5:F8)"
    ws["F9"].font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws["F9"].fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    ws["F9"].alignment = Alignment(horizontal="center", vertical="center")

    ws["G9"] = "=F9/B9"
    ws["G9"].number_format = "0.0%"
    ws["G9"].font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws["G9"].fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    ws["G9"].alignment = Alignment(horizontal="center", vertical="center")

    ws["H9"] = "=SUM(H5:H8)"
    ws["H9"].font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws["H9"].fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    ws["H9"].alignment = Alignment(horizontal="center", vertical="center")

    ws["I9"] = "=H9/B9"
    ws["I9"].number_format = "0.0%"
    ws["I9"].font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws["I9"].fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    ws["I9"].alignment = Alignment(horizontal="center", vertical="center")

    ws["J9"] = "=SUM(J5:J8)"
    ws["J9"].font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws["J9"].fill = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
    ws["J9"].alignment = Alignment(horizontal="center", vertical="center")

    ws["K9"] = "=J9/B9"
    ws["K9"].number_format = "0.0%"
    ws["K9"].font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws["K9"].fill = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
    ws["K9"].alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[9].height = 22
    for col_letter in ["A","B","C","D","E","F","G","H","I","J","K"]:
        ws[f"{col_letter}9"].border = thick_bottom

    # Footnote
    ws.merge_cells("A10:K10")
    fn = ws["A10"]
    fn.value = "* Note: Column C represents the true 100.0% Ownership Distribution across all identified issues."
    fn.font = Font(name="Calibri", size=9, italic=True, color="595959")
    fn.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[10].height = 18

    # -------------------------------------------------------------
    # SECTION 2: OWNER REACH (ALL CASES - COMBINED)
    # -------------------------------------------------------------
    ws.merge_cells("A12:L12")
    s2_cell = ws["A12"]
    s2_cell.value = "OWNER: REACH (ALL CASES - COMBINED)"
    s2_cell.fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    s2_cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    s2_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[12].height = 22

    # Headers for Left and Right Breakdown
    ws.merge_cells("A14:E14")
    w_left = ws["A14"]
    w_left.value = "ROOT CAUSE BREAKDOWN (ALL REACH CASES)"
    w_left.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    w_left.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    w_left.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("G14:I14")
    w_right = ws["G14"]
    w_right.value = "ROOT CAUSE SUMMARY (ALL)"
    w_right.fill = PatternFill(start_color="7030A0", end_color="7030A0", fill_type="solid")
    w_right.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    w_right.alignment = Alignment(horizontal="center", vertical="center")

    # Column Subheaders
    r_cols = [
        ("A15", "Reason", "1F4E79"),
        ("B15", "Total", "1F4E79"),
        ("C15", "Sub reason", "1F4E79"),
        ("D15", "Cases", "1F4E79"),
        ("E15", "RC (Operational Defect)", "1F4E79"),
        ("G15", "Root Cause Name", "7030A0"),
        ("H15", "Count", "7030A0"),
        ("I15", "Percentage (of Reach)", "7030A0"),
    ]
    for c_id, txt, bg in r_cols:
        c = ws[c_id]
        c.value = txt
        c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
        c.font = Font(name="Calibri", size=9.5, bold=True, color=WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center")

    reach_rows = [
        ("Non-issue feedback", 2, "Positive comment", 1, "Score education", "Score education", 2, "=H16/6"),
        ("Non-issue feedback", 2, "No stated reason", 1, "Score education", "Backoffice operations (Contractor)", 2, "=H17/6"),
        ("Installation delay", 2, "Delay in dispatching", 2, "Backoffice operations (Contractor)", "Service education", 1, "=H18/6"),
        ("Field Handover & Guidance", 1, "Technical orientation", 1, "Service education", "Bad splicing", 1, "=H19/6"),
        ("Installation quality", 1, "Service interruption", 1, "Bad splicing", "", "", ""),
    ]

    for i, r in enumerate(reach_rows, start=16):
        reason, total, sub, cases, rc, rc_name, rc_cnt, rc_pct = r
        ws[f"A{i}"] = reason
        ws[f"B{i}"] = total
        ws[f"C{i}"] = sub
        ws[f"D{i}"] = cases
        ws[f"E{i}"] = rc
        
        if rc_name:
            ws[f"G{i}"] = rc_name
            ws[f"H{i}"] = rc_cnt
            ws[f"I{i}"] = rc_pct
            ws[f"I{i}"].number_format = "0.0%"

        ws.row_dimensions[i].height = 19
        for cl in ["A","B","C","D","E"]:
            ws[f"{cl}{i}"].border = thin_border
            ws[f"{cl}{i}"].alignment = Alignment(horizontal="center" if cl in ["B","D"] else "left", vertical="center")
        if rc_name:
            for cl in ["G","H","I"]:
                ws[f"{cl}{i}"].border = thin_border
                ws[f"{cl}{i}"].alignment = Alignment(horizontal="center" if cl in ["H","I"] else "left", vertical="center")

    # Column Widths
    col_widths = {
        "A": 24, "B": 14, "C": 26, "D": 14, "E": 32,
        "F": 16, "G": 32, "H": 14, "I": 22, "J": 14,
        "K": 16, "L": 14
    }
    for col_l, width in col_widths.items():
        ws.column_dimensions[col_l].width = width

    # Save to desktop and project folder
    paths = [
        r"C:\Users\LENOVO\Desktop\NPS_Analysis_and_Feedback_Fixed.xlsx",
        "NPS_Analysis_and_Feedback_Fixed.xlsx"
    ]
    for p in paths:
        try:
            wb.save(p)
            print(f"[OK] Saved: {p}")
        except Exception as e:
            print(f"[WARN] Error saving {p}: {e}")

if __name__ == "__main__":
    create_fixed_nps_workbook()
