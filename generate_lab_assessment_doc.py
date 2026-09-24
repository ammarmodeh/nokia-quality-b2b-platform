import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f"""
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="none"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="none"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    """
    tblPr.append(parse_xml(borders_xml))

def create_lab_assessment_docx(output_path):
    doc = Document()

    # Page Margins (0.6 in for professional compact layout)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Styles & Colors
    NAVY = RGBColor(30, 58, 138)       # #1E3A8A
    DARK_BLUE = RGBColor(15, 23, 42)   # #0F172A
    TEAL = RGBColor(13, 148, 136)      # #0D9488
    PURPLE = RGBColor(124, 58, 237)    # #7C3AED
    GRAY_TEXT = RGBColor(100, 116, 139)# #64748B

    # Header Title Banner
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("E2E FTTH PRACTICAL LABORATORY ASSESSMENT EVALUATION FORM")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(17)
    title_run.font.bold = True
    title_run.font.color.rgb = NAVY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Nokia Quality Management System • Controlled Laboratory Assessment Standard\nFull Lifecycle: FDB → BEP → OTO → ONT (ZTE / Nokia) → Tools Readiness → Customer Education → Quality Awareness")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(9.5)
    sub_run.font.italic = True
    sub_run.font.color.rgb = GRAY_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # SECTION 1: FIELD TECHNICIAN PROFILE & EVALUATION METADATA
    # -------------------------------------------------------------
    h1 = doc.add_paragraph()
    h1_run = h1.add_run("1. FIELD TECHNICIAN PROFILE & ASSESSMENT METADATA")
    h1_run.font.name = "Calibri"
    h1_run.font.size = Pt(11)
    h1_run.font.bold = True
    h1_run.font.color.rgb = NAVY
    h1.paragraph_format.space_after = Pt(3)

    t1 = doc.add_table(rows=4, cols=4)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1, "1E3A8A", "6")

    col_widths_1 = [Inches(1.8), Inches(2.2), Inches(1.8), Inches(2.2)]
    
    fields_1 = [
        [("Field Team / Tech Name:", True), ("", False), ("Team Code / ID:", True), ("", False)],
        [("Subcontractor / Company:", True), ("", False), ("Employment Type:", True), ("[  ] Direct Employee   [  ] Subcontractor", False)],
        [("Date of Birth (DoB):", True), ("DD / MM / YYYY", False), ("Technician Age:", True), ("_____ Years", False)],
        [("Experience Duration:", True), ("___ Yrs, ___ Mos, ___ Days", False), ("Assessor & Date:", True), ("Name: ______________  Date: ___/___/2026", False)],
    ]

    for r_idx, row in enumerate(t1.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths_1[c_idx]
            text, is_label = fields_1[r_idx][c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(9)
            if is_label:
                run.font.bold = True
                set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, 60, 60, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 2: SAFETY & COMPLIANCE CLEARANCE MATRIX (6 CERTS)
    # -------------------------------------------------------------
    h2 = doc.add_paragraph()
    h2_run = h2.add_run("2. MANDATORY SAFETY & COMPLIANCE CLEARANCE MATRIX (6 CORE CERTIFICATES)")
    h2_run.font.name = "Calibri"
    h2_run.font.size = Pt(11)
    h2_run.font.bold = True
    h2_run.font.color.rgb = PURPLE
    h2.paragraph_format.space_after = Pt(3)

    t2 = doc.add_table(rows=7, cols=4)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2, "7C3AED", "6")

    col_widths_2 = [Inches(1.0), Inches(2.2), Inches(3.4), Inches(1.4)]
    
    headers_2 = ["Code", "Certificate Title", "Safety Scope & Mandatory Verification", "Verification Status"]
    for c_idx, heading in enumerate(headers_2):
        cell = t2.rows[0].cells[c_idx]
        cell.width = col_widths_2[c_idx]
        set_cell_background(cell, "7C3AED")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(heading)
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 80, 80, 80, 80)

    cert_data = [
        ("DD", "Defensive Driving", "Safe vehicle operation, hazard perception, seatbelt & driving compliance", "[  ] Valid   [  ] Missing"),
        ("ES", "Electrical Safety", "Safe handling of electrical sources, grounding & arc flash prevention", "[  ] Valid   [  ] Missing"),
        ("FA", "First Aid", "Emergency first response, CPR & basic field trauma management", "[  ] Valid   [  ] Missing"),
        ("FF", "Fire Fighting", "Fire extinguisher operation, fire risk prevention & emergency evacuation", "[  ] Valid   [  ] Missing"),
        ("RA", "Risk Assessment", "Job site hazard identification, risk evaluation & dynamic control measures", "[  ] Valid   [  ] Missing"),
        ("WAH", "Work at Height", "Ladder safety, harness usage, pole climbing & fall arrest systems", "[  ] Valid   [  ] Missing"),
    ]

    for r_idx, (code, title, desc, status) in enumerate(cert_data):
        row = t2.rows[r_idx + 1]
        for c_idx, val in enumerate([code, title, desc, status]):
            cell = row.cells[c_idx]
            cell.width = col_widths_2[c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            run = p.add_run(val)
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run.font.bold = True
                set_cell_background(cell, "F3E8FF")
            elif c_idx == 1:
                run.font.bold = True
            elif c_idx == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_margins(cell, 50, 50, 80, 80)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 3: Tools & Equipment
    # -------------------------------------------------------------
    h3 = doc.add_paragraph()
    h3_run = h3.add_run("3. Tools & Equipment")
    h3_run.font.name = "Calibri"
    h3_run.font.size = Pt(11)
    h3_run.font.bold = True
    h3_run.font.color.rgb = RGBColor(234, 88, 12) # Orange
    h3.paragraph_format.space_after = Pt(3)

    t3 = doc.add_table(rows=12, cols=4)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t3, "EA580C", "6")

    col_widths_3 = [Inches(0.6), Inches(2.6), Inches(2.8), Inches(2.0)]
    
    headers_3 = ["Item #", "Tool / Equipment Item Name", "Inspection Status Check", "Assessor Condition / Serial Remarks"]
    for c_idx, heading in enumerate(headers_3):
        cell = t3.rows[0].cells[c_idx]
        cell.width = col_widths_3[c_idx]
        set_cell_background(cell, "EA580C")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(heading)
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 80, 80, 80, 80)

    tools_data = [
        ("1", "Splicing Machine", "[  ] Valid   [  ] Need Maint.   [  ] Not Valid   [  ] Missing"),
        ("2", "Fiber Cleaver", "[  ] Valid   [  ] Need Maint.   [  ] Not Valid   [  ] Missing"),
        ("3", "Fiber Stripper", "[  ] Valid   [  ] Need Maint.   [  ] Not Valid   [  ] Missing"),
        ("4", "Optical Power Meter (OPM)", "[  ] Valid   [  ] Need Maint.   [  ] Not Valid   [  ] Missing"),
        ("5", "Laser Source (VFL)", "[  ] Valid   [  ] Need Maint.   [  ] Not Valid   [  ] Missing"),
        ("6", "Cleaning Solution – IPA 99%", "[  ] Available   [  ] Not Available"),
        ("7", "Lint-Free Cotton Swabs / Ear Sticks", "[  ] Available   [  ] Not Available"),
        ("8", "Cleaning Wipes – Kimwipes", "[  ] Available   [  ] Not Available"),
        ("9", "Cleaning Brush", "[  ] Available   [  ] Not Available"),
        ("10", "Protection Sleeves (BEP / FDB / OTO)", "[  ] Available   [  ] Not Available"),
        ("11", "Labels (FDB / BEP / OTO / ONT)", "[  ] Available   [  ] Not Available"),
    ]

    for r_idx, (num, tool, status_opts) in enumerate(tools_data):
        row = t3.rows[r_idx + 1]
        for c_idx, val in enumerate([num, tool, status_opts, ""]):
            cell = row.cells[c_idx]
            cell.width = col_widths_3[c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            run = p.add_run(val)
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run.font.bold = True
            elif c_idx == 1:
                run.font.bold = True
            set_cell_margins(cell, 40, 40, 60, 60)

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 4: ITEMIZED PRACTICAL LAB ASSESSMENT CHECKPOINTS
    # -------------------------------------------------------------
    h4 = doc.add_paragraph()
    h4_run = h4.add_run("4. ITEMIZED PRACTICAL LABORATORY CHECKPOINTS (SECTIONS 1 TO 8)")
    h4_run.font.name = "Calibri"
    h4_run.font.size = Pt(13)
    h4_run.font.bold = True
    h4_run.font.color.rgb = NAVY
    h4.paragraph_format.space_after = Pt(2)

    guide_p = doc.add_paragraph()
    guide_run = guide_p.add_run("Scoring Strategy: Checkpoints in Sections 1–7 are rated from 0 to 5 Points (0, 1, 2, 3 = Fail | 4, 5 = Pass). Section 8 is rated from 0 to 10 Points (0–6 = Fail | 7–10 = Pass). Section 7 is Informational Quality Awareness.")
    guide_run.font.name = "Calibri"
    guide_run.font.size = Pt(8.5)
    guide_run.font.italic = True
    guide_run.font.color.rgb = DARK_BLUE
    guide_p.paragraph_format.space_after = Pt(4)

    # Detailed Checklist Definition
    CHECKLIST_SECTIONS = [
        {
            "id": "1",
            "title": "1. FDB – Fiber Splicing & Fiber Management",
            "color": "1E3A8A",
            "checkpoints": [
                ("A. Power Measurement & Splitter Verification", "Measure optical power received from selected splitter port before starting installation", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("A. Power Measurement & Splitter Verification", "Record measured optical power and determine whether within theoretical acceptable range (-16 to -19 dBm)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("A. Power Measurement & Splitter Verification", "Record measured optical power and determine whether within practical acceptable range (-16 to -22 or -23 dBm)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. New Drop Fiber Installation", "Properly perform fusion splicing between pigtail and outdoor drop fiber (black)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. New Drop Fiber Installation", "Properly route and organize pigtail inside cassette/raceway without excessive bending, tension, or crossing", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. New Drop Fiber Installation", "Properly route and organize outdoor drop fiber inside cassette/raceway without excessive bending, tension, or crossing", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. New Drop Fiber Installation", "Install required label on Pigtail (Splitter port / identification label)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. New Drop Fiber Installation", "Install required label on Outdoor drop fiber cable (black)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. New Drop Fiber Installation", "Ensure splice is properly protected with sleeve and securely positioned", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("C. Existing Drop Fiber Handling", "Identify required existing drop fiber and carefully extract/reposition from raceway without disturbing others", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("C. Existing Drop Fiber Handling", "Maintain proper fiber organization, avoid cassette deformation, and return fibers to correct position", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
            ]
        },
        {
            "id": "2",
            "title": "2. BEP – Outdoor/Indoor Fiber Splicing & Fiber Management",
            "color": "0D9488",
            "checkpoints": [
                ("BEP Splicing & Management", "Perform fusion splicing between Outdoor drop fiber (black) and Indoor drop fiber (white)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("BEP Splicing & Management", "Properly organize and route fiber inside BEP according to required installation standard", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("BEP Splicing & Management", "Avoid excessive bending, tension, or unnecessary fiber loops, and properly secure fiber/cable", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("BEP Labeling", "Apply required label on Outdoor drop fiber (black) at BEP (clear, readable, correct position)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("BEP Labeling", "Apply required label on Indoor drop fiber (white) at BEP (clear, readable, correct position)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
            ]
        },
        {
            "id": "3",
            "title": "3. OTO – Splicing, Fiber Management & Optical Testing",
            "color": "B45309",
            "checkpoints": [
                ("A. Splicing", "Perform fusion splicing between indoor drop fiber and built-in OTO pigtail", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("A. Splicing", "Properly protect and position splice, organize fiber inside OTO without excessive bending or tension", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. Labeling", "Apply required identification label to OTO cover (correctly positioned and readable)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("C. Optical Testing", "Connect patch cord from OTO to optical power meter and perform required optical test", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("C. Optical Testing", "Measure link loss, record value, and verify within acceptable range (at max 2 dB difference from FDB splitter Rx power)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
            ]
        },
        {
            "id": "4",
            "title": "4. ONT – Installation, Configuration & Service Activation (ZTE & Nokia Models)",
            "color": "9D174D",
            "checkpoints": [
                ("A. ZTE F6600P Configuration", "ZTE F6600P – Power on ONT and enter required INT number in the appropriate field", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("A. ZTE F6600P Configuration", "ZTE F6600P – Configure required WAN connection information correctly", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("A. ZTE F6600P Configuration", "ZTE F6600P – Configure Wi-Fi settings separately: 2.4 GHz & 5 GHz (SSIDs & Passwords)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("A. ZTE F6600P Configuration", "ZTE F6600P – Check & verify received optical signal level", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("A. ZTE F6600P Service Verification", "ZTE F6600P – Understand and explain ONT LED behavior & status indicators (PON solid green, LOS red, WAN & Wi-Fi LEDs)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. Nokia G-2426G-P (Wi-Fi 6)", "Nokia G-2426G-P – Power on ONT and enter required INT number in the appropriate field", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. Nokia G-2426G-P (Wi-Fi 6)", "Nokia G-2426G-P – Mandatory: Manually upload firmware version 62 and verify automatic restart before continuing", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. Nokia G-2426G-P (Wi-Fi 6)", "Nokia G-2426G-P – Configure required WAN connection information correctly", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. Nokia G-2426G-P (Wi-Fi 6)", "Nokia G-2426G-P – Configure Wi-Fi settings separately: 2.4 GHz & 5 GHz (SSIDs & Passwords)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. Nokia G-2426G-P (Wi-Fi 6)", "Nokia G-2426G-P – Check & verify received optical signal level", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. Nokia G-2426G-P (Wi-Fi 6)", "Nokia G-2426G-P – Understand and explain ONT LED behavior & status indicators (PON solid green, LOS red, WAN & Wi-Fi LEDs)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("C. Nokia G-140W-H (Wi-Fi 5)", "Nokia G-140W-H – Configure required WAN connection information correctly", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("C. Nokia G-140W-H (Wi-Fi 5)", "Nokia G-140W-H – Configure Wi-Fi settings separately: 2.4 GHz & 5 GHz (SSIDs & Passwords)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("C. Nokia G-140W-H (Wi-Fi 5)", "Nokia G-140W-H – Check & verify received optical signal level", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("C. Nokia G-140W-H (Wi-Fi 5)", "Nokia G-140W-H – Understand and explain ONT LED behavior & status indicators (PON solid green, LOS red, WAN & Wi-Fi LEDs)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("D. ONT Labeling (Shared)", "Apply required Call Center & INT number label on ONT (clearly visible to customer)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("D. ONT Labeling (Shared)", "Apply required NO RESET label on ONT (placed directly over/near reset-hole area, clearly visible to customer)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
            ]
        },
        {
            "id": "5",
            "title": "5. Assistant Tools & Mobile Applications Verification",
            "color": "047857",
            "checkpoints": [
                ("A. Ookla Speedtest App", "Verify team smartphone has Ookla Speedtest application installed and properly configured (server selection & testing standard)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. Wi-Fi Signal Meter", "Verify team smartphone has Wi-Fi Signal Meter application installed and configured to measure 2.4 GHz & 5 GHz coverage", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("C. Timestamp Camera App", "Verify team smartphone has Timestamp Camera application installed with correct date, time, and GPS coordinates configuration", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("D. 1 Gbps-Capable Laptop", "1 Gbps-capable laptop: Team must have a laptop with 1 Gbps+ Ethernet interface for high-speed internet speed verification", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("E. High-Speed Wi-Fi Device", "High-speed Wi-Fi device: Team must have a Wi-Fi 6–capable laptop or mobile device to verify high-speed Wi-Fi 6 router performance", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
            ]
        },
        {
            "id": "6",
            "title": "6. Customer Education & Communication Protocols",
            "color": "4338CA",
            "checkpoints": [
                ("A. Wi-Fi Education", "Explain difference between 2.4 GHz and 5 GHz Wi-Fi (advantages, limitations, distance vs speed)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("B. 3-Point Speed Test", "Explain and demonstrate/perform required 3-point speed test per project procedure", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("C. ONT / Service Handling", "Clearly explain NOT to reset ONT and that pressing reset removes configuration & interrupts service", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("D. Support Routing", "Educate customer on support routing: Call Center (Sales/Billing/Accounts) vs Tech Team (Internet/Speed/Issues)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("E. CX Follow-Up Call", "Inform customer of CX follow-up call after ~48h and explain CSAT satisfaction scale (1–8 Unsatisfied, 9–10 Satisfied)", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
                ("F. Service Education", "Educate customer on expected performance for IPTV, VPN, PlayStation/Sony servers, and server speed differences", "0 - 5 Pts", "[  ] P   [  ] F   [  ] N/A"),
            ]
        },
        {
            "id": "7",
            "title": "7. Quality Ticket Awareness (Informational Standard)",
            "color": "6D28D9",
            "checkpoints": [
                ("Ticket Types & SLAs", "Demonstrate clear understanding of NPS Ticket (SLA: 24 hours) and CC Ticket (SLA: 9 hours)", "0 - 5 Pts", "[  ] Pass   [  ] Fail"),
                ("SLA Compliance & Escalation", "Demonstrate understanding of timely corrective action within SLA and proper escalation when SLA is at risk", "0 - 5 Pts", "[  ] Pass   [  ] Fail"),
            ]
        },
        {
            "id": "8",
            "title": "8. Self-Assessment & Photo Documentation (10 Points)",
            "color": "0284C7",
            "checkpoints": [
                ("Step-by-Step Installation Photo Verification & Quality Self-Assessment", "Capture complete, clear Timestamp photo evidence for every step (FDB splicing & labels, BEP splicing & labels, OTO splicing & cover label, ONT installation & labels, Speedtest result, and Wi-Fi signal meter) to thoroughly self-assess installation quality", "0 - 10 Pts (Score: ___ / 10)", "[  ] Pass (≥7)   [  ] Fail (<7)"),
            ]
        },
    ]

    col_widths_chk = [Inches(0.4), Inches(2.0), Inches(3.2), Inches(1.0), Inches(1.4)]

    for sec in CHECKLIST_SECTIONS:
        # Section Header Paragraph
        sec_p = doc.add_paragraph()
        sec_run = sec_p.add_run(sec["title"])
        sec_run.font.name = "Calibri"
        sec_run.font.size = Pt(10.5)
        sec_run.font.bold = True
        sec_run.font.color.rgb = RGBColor.from_string(sec["color"])
        sec_p.paragraph_format.space_before = Pt(8)
        sec_p.paragraph_format.space_after = Pt(2)

        # Table
        t_sec = doc.add_table(rows=len(sec["checkpoints"]) + 1, cols=5)
        t_sec.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t_sec, sec["color"], "6")

        # Table Header
        h_row = t_sec.rows[0]
        h_titles = ["#", "Sub-Section / Scope", "Checkpoint & Technical Criteria", "Points", "Result & Observations"]
        for c_idx, title in enumerate(h_titles):
            cell = h_row.cells[c_idx]
            cell.width = col_widths_chk[c_idx]
            set_cell_background(cell, sec["color"])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(title)
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            set_cell_margins(cell, 60, 60, 60, 60)

        for cp_idx, (sub_sec, criteria, pts, res) in enumerate(sec["checkpoints"]):
            row = t_sec.rows[cp_idx + 1]
            vals = [str(cp_idx + 1), sub_sec, criteria, pts, res]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                cell.width = col_widths_chk[c_idx]
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                run = p.add_run(val)
                run.font.name = "Calibri"
                run.font.size = Pt(8)
                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run.font.bold = True
                elif c_idx == 1:
                    run.font.bold = True
                elif c_idx == 3:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run.font.bold = True
                elif c_idx == 4:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_margins(cell, 40, 40, 60, 60)

    doc.add_page_break()

    # -------------------------------------------------------------
    # SECTION 5: SCORING BENCHMARK SCALE & EVALUATION SUMMARY
    # -------------------------------------------------------------
    h5 = doc.add_paragraph()
    h5_run = h5.add_run("5. SCORING BENCHMARK SCALE & PERFORMANCE EVALUATION SUMMARY")
    h5_run.font.name = "Calibri"
    h5_run.font.size = Pt(12)
    h5_run.font.bold = True
    h5_run.font.color.rgb = NAVY
    h5.paragraph_format.space_after = Pt(3)

    t5 = doc.add_table(rows=6, cols=4)
    t5.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t5, "1E3A8A", "6")

    col_widths_5 = [Inches(1.5), Inches(1.2), Inches(3.5), Inches(1.8)]
    headers_5 = ["Score Range (%)", "Evaluation Status", "5-Tier Quality Performance Classification", "Operational Action Standard"]
    
    for c_idx, title in enumerate(headers_5):
        cell = t5.rows[0].cells[c_idx]
        cell.width = col_widths_5[c_idx]
        set_cell_background(cell, "1E3A8A")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 60, 60, 60, 60)

    scale_data = [
        ("90.0% – 100.0%", "Pass", "Excellent", "Authorized for live FTTH subscriber deployments", "DCFCE7"),
        ("80.0% – 89.9%", "Pass", "V.Good", "Authorized for standard FTTH field installations", "ECFDF5"),
        ("70.0% – 79.9%", "Pass", "Need Improvements & Close Monitoring", "Conditional pass: requires targeted coaching", "FEF9C3"),
        ("50.0% – 69.9%", "Fail", "Suspension, Need Training & Re-evaluation", "Work suspended: mandatory lab retraining & retest", "FFEDD5"),
        ("0.0% – 49.9%", "Fail", "Ineligible", "Unauthorized for field work: fundamental gap", "FEE2E2"),
    ]

    for r_idx, (rng, stat, tier, action, bg_hex) in enumerate(scale_data):
        row = t5.rows[r_idx + 1]
        for c_idx, val in enumerate([rng, stat, tier, action]):
            cell = row.cells[c_idx]
            cell.width = col_widths_5[c_idx]
            set_cell_background(cell, bg_hex)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            run = p.add_run(val)
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            if c_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run.font.bold = True
            elif c_idx == 2:
                run.font.bold = True
            set_cell_margins(cell, 40, 40, 60, 60)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # SECTION 6: ASSESSOR COMPREHENSIVE FEEDBACK & CLEARANCE SIGN-OFF
    # -------------------------------------------------------------
    h6 = doc.add_paragraph()
    h6_run = h6.add_run("6. ASSESSOR COMPREHENSIVE FEEDBACK, ACTION PLAN & CLEARANCE SIGN-OFF")
    h6_run.font.name = "Calibri"
    h6_run.font.size = Pt(12)
    h6_run.font.bold = True
    h6_run.font.color.rgb = NAVY
    h6.paragraph_format.space_after = Pt(3)

    t6 = doc.add_table(rows=2, cols=1)
    t6.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t6, "1E3A8A", "6")

    cell_feedback = t6.rows[0].cells[0]
    cell_feedback.width = Inches(8.0)
    set_cell_background(cell_feedback, "F8FAFC")
    p_f = cell_feedback.paragraphs[0]
    p_f.add_run("Assessor Observations, Strengths & Specific Technical Corrective Action Plan:\n\n\n\n\n").font.size = Pt(9)
    set_cell_margins(cell_feedback, 60, 60, 80, 80)

    t7 = doc.add_table(rows=2, cols=3)
    t7.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t7, "1E3A8A", "6")

    widths_7 = [Inches(2.6), Inches(2.6), Inches(2.8)]
    sign_headers = ["Field Technician Acknowledgement", "Quality Assessor Sign-off", "Quality Supervisor Clearance"]
    for c_idx, heading in enumerate(sign_headers):
        cell = t7.rows[0].cells[c_idx]
        cell.width = widths_7[c_idx]
        set_cell_background(cell, "1E3A8A")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(heading)
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 60, 60, 60, 60)

    sign_body = [
        "Technician Name: __________________\nSignature: ________________________\nDate: _____ / _____ / 2026",
        "Assessor Name: ____________________\nSignature: ________________________\nDate: _____ / _____ / 2026",
        "Supervisor Name: __________________\nDecision: [  ] Approved   [  ] Retest\nSignature: ________________________",
    ]

    for c_idx, body in enumerate(sign_body):
        cell = t7.rows[1].cells[c_idx]
        cell.width = widths_7[c_idx]
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(body)
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        set_cell_margins(cell, 60, 60, 60, 60)

    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    out_dir = r"c:\Users\LENOVO\Desktop\Coding\B1 MERN Projects\32 FullstackMERN Stack Task Manager App - nokia-quality-task-tracker\Project"
    out_file = os.path.join(out_dir, "Nokia_E2E_FTTH_Practical_Laboratory_Assessment_Form.docx")
    create_lab_assessment_docx(out_file)
